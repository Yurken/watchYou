"""Build a styled DOCX from the project's Markdown PRD. Requires python-docx."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'docs/watch_you_PRD.md'
OUTPUT=ROOT/'docs/watch_you_PRD.docx'
FONT='Noto Sans CJK SC'
INK='223B4C'; BLUE='42677E'; MUTED='627B8A'; PALE='ECF2F5'

def fonts(obj,name=FONT,size=None,color=None,bold=None):
    obj.font.name=name
    if size is not None: obj.font.size=Pt(size)
    if color: obj.font.color.rgb=RGBColor.from_string(color)
    if bold is not None:obj.font.bold=bold
    rpr=obj.element.get_or_add_rPr() if hasattr(obj,'element') else obj._element.get_or_add_rPr()
    rf=rpr.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');rpr.insert(0,rf)
    for x in ('ascii','hAnsi','eastAsia','cs'):rf.set(qn('w:'+x),name)

def field(p,code):
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),code)
    rr=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1';rr.append(t);f.append(rr);p._p.append(f)

def shade(cell,fill):
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),fill);cell._tc.get_or_add_tcPr().append(sh)

def inline(p,text,size=None,color=None):
    for v in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)',text):
        if not v:continue
        isbold=v.startswith('**') and v.endswith('**')
        iscode=v.startswith('`') and v.endswith('`')
        if isbold:v=v[2:-2]
        if iscode:v=v[1:-1]
        rr=p.add_run(v)
        fonts(rr, 'DejaVu Sans Mono' if iscode and v.isascii() else FONT,size=size,color=color,bold=isbold)

def table_style(t,widths):
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+side);x.set(qn('w:val'),'single' if side in ['top','bottom','insideH'] else 'nil');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D7E3E9');borders.append(x)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for k,n in [('top',75),('bottom',75),('left',90),('right',90)]:
        x=OxmlElement('w:'+k);x.set(qn('w:w'),str(n));x.set(qn('w:type'),'dxa');margins.append(x)
    pr.append(margins)
    for col,w in zip(t.columns,widths):col.width=Mm(w)
    for row in t.rows:
        np=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(np)
        for i,c in enumerate(row.cells):c.width=Mm(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def md_table(rows):
    n=len(rows[0]);t=doc.add_table(rows=len(rows),cols=n)
    ratios={2:[.28,.72],3:[.21,.32,.47],4:[.19,.27,.31,.23]}.get(n,[1/n]*n)
    widths=[170*x for x in ratios]
    if n==3 and rows[0][0]=='编号':widths=[15,58,97]
    if n==3 and rows[0][0]=='字段':widths=[43,54,73]
    if n==3 and rows[0][0]=='实体':widths=[34,76,60]
    table_style(t,widths)
    for i,row in enumerate(rows):
        for j,v in enumerate(row):
            c=t.cell(i,j);p=c.paragraphs[0]
            p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
            p.paragraph_format.line_spacing=1.16
            inline(p,v,size=9,color=INK)
            if i==0:
                shade(c,PALE)
                for r in p.runs:r.bold=True
            elif i%2==0:shade(c,'F8FAFB')
        if i==0:
            repeat=OxmlElement('w:tblHeader');t.rows[i]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

text=SOURCE.read_text(encoding='utf-8')
doc=Document();sec=doc.sections[0]
sec.page_width=Mm(210);sec.page_height=Mm(297)
sec.left_margin=Mm(20);sec.right_margin=Mm(20);sec.top_margin=Mm(19);sec.bottom_margin=Mm(18)
sec.header_distance=Mm(8);sec.footer_distance=Mm(8);sec.different_first_page_header_footer=True
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    st=doc.styles[name];fonts(st,size=10.5,color=INK)
normal=doc.styles['Normal'];normal.paragraph_format.line_spacing=1.22;normal.paragraph_format.space_after=Pt(6)
for name,size,before,after in [('Heading 1',16,16,9),('Heading 2',12,11,6),('Heading 3',11,8,4)]:
    st=doc.styles[name];fonts(st,size=size,color=BLUE,bold=True)
    st.paragraph_format.space_before=Pt(before);st.paragraph_format.space_after=Pt(after);st.paragraph_format.keep_with_next=True
    st.paragraph_format.keep_together=True
fonts(doc.styles['Caption'],size=8.5,color=MUTED)
# Header and footer.
h=sec.header.paragraphs[0];h.paragraph_format.space_after=Pt(0)
fonts(h.add_run('watch you'),size=9,bold=True,color=BLUE)
h.add_run('    /    AI HEALTH COMPANION · PRODUCT REQUIREMENTS').font.size=Pt(8)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fonts(f.add_run('watch you · v2.0    |    '),size=8,color=MUTED);field(f,'PAGE')
fonts(f.add_run(' / '),size=8,color=MUTED);field(f,'NUMPAGES')
# Cover.
p=doc.add_paragraph();p.paragraph_format.space_before=Pt(16);p.paragraph_format.space_after=Pt(6)
fonts(p.add_run('watch you'),size=38,bold=True,color=INK)
p=doc.add_paragraph();fonts(p.add_run('AI 健康与情绪助手'),size=20,color=BLUE,bold=True)
p.paragraph_format.space_after=Pt(6)
p=doc.add_paragraph();fonts(p.add_run('完整产品需求文档  /  iPhone + Apple Watch'),size=12,color=MUTED)
p=doc.add_paragraph();fonts(p.add_run('v2.0  ·  2026-10-21  ·  Yurken/watchYou'),size=10,color=MUTED)
p.paragraph_format.space_after=Pt(12)
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(ROOT/'assets/desktop-overview.png'),width=Mm(170))
# Set image alt text without external font assets.
for shape in doc.inline_shapes:
    shape._inline.docPr.set('descr','watch you 手机与手表交互原型；所有健康数值均为合成示例。')
p=doc.add_paragraph('自动情绪洞察 · 全模块 AI · 自定义模型服务 · 通用数据来源')
fonts(p.runs[0],size=10.5,bold=True,color=BLUE)
p=doc.add_paragraph('交付状态：需求与浏览器交互原型。尚未完成原生实机接入。')
fonts(p.runs[0],size=9,color=MUTED)
p=doc.add_paragraph('原型全部健康数据为合成示例；AI 推断不等同于本人自述或医学事实。真实模型调用不代表真实 HealthKit 接入。')
fonts(p.runs[0],size=9,color=MUTED)
doc.add_page_break()
# Static chapter navigation; headings also populate Word's navigation pane.
p=doc.add_paragraph('阅读导航',style='Heading 1')
p=doc.add_paragraph('正文共 28 章。先阅读产品与 AI 决策，再进入技术规范和验收。Word 导航窗格可按章节跳转。')
fonts(p.runs[0],size=10,color=MUTED)
chapters=re.findall(r'^## (.+)$',text,re.M)
t=doc.add_table(rows=14,cols=2);table_style(t,[85,85])
for i in range(14):
    for j,k in enumerate([i,i+14]):
        c=t.cell(i,j);p=c.paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
        fonts(p.add_run(chapters[k]),size=9.8,color=INK)
        if i%2==0:shade(c,'F3F7F9')
p=doc.add_paragraph('本版重点',style='Heading 2')
for v in ['第 08 章：自动情绪推断与可选反馈。','第 11、25–27 章：全模块 AI、服务配置、契约与自动调度。','第 13 章：来源治理不依赖设备数量；第 18 章保留每周维护要求。','第 22–23 章区分产品验收目标与本次实际验证；第 28 章说明仓库交接。']:
    p=doc.add_paragraph(v);fonts(p.runs[0],size=10,color=MUTED)
doc.add_page_break()
# Actual Markdown body.
lines=text[text.index('## 01 '):].splitlines();i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('## '):doc.add_paragraph(line[3:],style='Heading 1');i+=1;continue
    if line.startswith('### '):doc.add_paragraph(line[4:],style='Heading 2');i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            vals=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',v) for v in vals):rows.append(vals)
            i+=1
        md_table(rows);continue
    if line.startswith('```'):
        i+=1
        while i<len(lines) and not lines[i].startswith('```'):
            p=doc.add_paragraph();fonts(p.add_run(lines[i]),name='DejaVu Sans Mono',size=8.5,color=INK);i+=1
        i+=1;continue
    p=doc.add_paragraph()
    if line.startswith('http'):
        p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=1.12
        inline(p,line,size=8.5,color=BLUE)
    else:
        inline(p,line)
        if re.match(r'^\[S\d+\]',line):
            p.paragraph_format.keep_with_next=True
            for r in p.runs:r.font.size=Pt(9.3)
    i+=1
props=doc.core_properties;props.title='watch you · AI 健康与情绪助手 · 完整产品需求文档';props.subject='v2.0 产品需求、AI 服务设置、数据治理、交互与验收';props.author='watch you';props.keywords='watch you, HealthKit, watchOS, AI, PRD';props.comments='设计与浏览器原型，不是已验证的原生或医疗产品。'
doc.save(OUTPUT)
print(OUTPUT)
